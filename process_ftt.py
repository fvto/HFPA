import copy as pycopy
import glob
import os
import re
import sys
import warnings
import argparse

warnings.filterwarnings("ignore", category=UserWarning)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="backslashreplace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="backslashreplace")

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import (
    RichTextProperties,
    Paragraph,
    ParagraphProperties,
    CharacterProperties,
)
import pandas as pd

from pipeline_common import extract_site_name, save_with_fallback

# Define paths relative to the current script directory
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_FOLDER = os.path.join(BASE_DIR, "FTT_Input")
OUTPUT_FOLDER = os.path.join(BASE_DIR, "Output")
OUTPUT_FILE = os.path.join(OUTPUT_FOLDER, "FTT_Combined_Report.xlsx")

# Defect -> color mapping workbook (Defect Name | Color | HEX | R | G | B)
COLOR_MAP_FILE = os.path.join(BASE_DIR, "BC_Color", "Color_Defect.xlsx")

# CoPQ cost workbook. The "Current month" sheet contains per-site blocks
# (headers: FTY | Model | Defect Qty | Ttl cost ($)) used to rank the Top 3
# Models per site by total cost instead of by defect quantity.
COST_FILE = os.path.join(BASE_DIR, "Database", "CoPQ_type_analysis.xlsx")
COST_SHEET_NAME = "Current month"

# Fallback palette used only if a defect isn't found in the color map file,
# so every series still gets a distinct, readable color.
FALLBACK_PALETTE = [
    "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5",
    "70AD47", "264478", "9E480E", "636363", "997300",
    "255E91", "43682B", "698ED0", "F1975A", "B7B7B7",
]


# ----------------------------------------------------------------------
# 1b. HELPER FUNCTIONS: DEFECT COLOR MAPPING & CHART STYLING
# ----------------------------------------------------------------------
def _normalize_defect_name(name):
    """Normalize a defect name for matching: lowercase, no whitespace."""
    return re.sub(r"\s+", "", str(name).strip().lower())


def _to_hex_str(value):
    """Best-effort conversion of a HEX cell's value to a clean 'RRGGBB' string."""
    if value is None:
        return None
    s = str(value).strip()
    if s.startswith("#"):
        s = s[1:]
    if re.fullmatch(r"[0-9A-Fa-f]{6}", s):
        return s.upper()
    try:
        num = int(float(value))
        s2 = str(abs(num))
        if 1 <= len(s2) <= 6:
            return s2.zfill(6).upper()
    except (TypeError, ValueError):
        pass
    return None


def _rgb_to_hex(r_val, g_val, b_val):
    """Convert R/G/B cell values to 'RRGGBB', clamping out-of-range values."""
    try:
        r = max(0, min(255, int(float(r_val))))
        g = max(0, min(255, int(float(g_val))))
        b = max(0, min(255, int(float(b_val))))
        return f"{r:02X}{g:02X}{b:02X}"
    except (TypeError, ValueError):
        return None


def _find_color_table_header(ws, max_scan_rows=15):
    """Scan the first few rows of a sheet to locate header columns for defect name, HEX, R/G/B."""
    max_col = ws.max_column or 1
    for r in range(1, max_scan_rows + 1):
        name_col = hex_col = r_col = g_col = b_col = None
        for c in range(1, max_col + 1):
            val = ws.cell(row=r, column=c).value
            if val is None:
                continue
            h = str(val).strip().lower()
            if not h:
                continue
            if name_col is None and "defect" in h:
                name_col = c
            elif hex_col is None and h == "hex":
                hex_col = c
            elif r_col is None and h in ("red", "r"):
                r_col = c
            elif g_col is None and h in ("green", "g"):
                g_col = c
            elif b_col is None and h in ("blue", "b"):
                b_col = c
        if name_col and (hex_col or (r_col and g_col and b_col)):
            return r, name_col, hex_col, r_col, g_col, b_col
    return None


def load_defect_color_map(source, sheet_names=None):
    """Load the defect -> HEX color mapping from workbook or file."""
    color_map = {}

    if isinstance(source, str):
        path = source
        if not os.path.exists(path):
            print(f"[!] Color mapping file not found: {path}")
            print("    Charts will use automatic fallback color palette.")
            return color_map
        try:
            wb_src = openpyxl.load_workbook(path, data_only=True)
        except Exception as e:
            print(f"[!] Could not read color mapping file ({path}): {e}")
            return color_map
        sheets_to_scan = wb_src.sheetnames
        source_label = path
    else:
        wb_src = source
        sheets_to_scan = sheet_names or wb_src.sheetnames
        source_label = "embedded workbook sheets"

    total_loaded = 0
    for sheet_name in sheets_to_scan:
        if sheet_name not in wb_src.sheetnames:
            continue
        ws = wb_src[sheet_name]
        header = _find_color_table_header(ws)
        if not header:
            continue
        header_row, name_col, hex_col, r_col, g_col, b_col = header

        sheet_loaded = 0
        for r in range(header_row + 1, ws.max_row + 1):
            defect_name = ws.cell(row=r, column=name_col).value
            if not defect_name or not str(defect_name).strip():
                continue

            hex_val = _to_hex_str(ws.cell(row=r, column=hex_col).value) if hex_col else None
            if not hex_val and r_col and g_col and b_col:
                hex_val = _rgb_to_hex(
                    ws.cell(row=r, column=r_col).value,
                    ws.cell(row=r, column=g_col).value,
                    ws.cell(row=r, column=b_col).value,
                )

            if not hex_val:
                continue

            color_map[_normalize_defect_name(defect_name)] = hex_val
            sheet_loaded += 1

        print(f"  [+] Loaded {sheet_loaded} defect colors from sheet '{sheet_name}'")
        total_loaded += sheet_loaded

    print(f"  [+] Total defect colors loaded: {len(color_map)} (from: {source_label})")
    return color_map


def embed_color_defect_workbook(wb, color_file_path):
    """Copy sheets from Color_Defect.xlsx into output report workbook."""
    if not os.path.exists(color_file_path):
        print(f"[!] Color mapping file not found, cannot embed: {color_file_path}")
        return []

    try:
        src_wb = openpyxl.load_workbook(color_file_path, data_only=True)
    except Exception as e:
        print(f"[!] Could not read color mapping file ({color_file_path}): {e}")
        return []

    added_sheets = []
    for sheet_name in src_wb.sheetnames:
        src_ws = src_wb[sheet_name]
        base_name = f"ColorMap_{sheet_name}"[:31]
        new_name = base_name
        suffix = 1
        while new_name in wb.sheetnames:
            suffix += 1
            new_name = f"{base_name[:28]}_{suffix}"

        dst_ws = wb.create_sheet(title=new_name)

        for row in src_ws.iter_rows():
            for cell in row:
                new_cell = dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    new_cell.font = pycopy.copy(cell.font)
                    new_cell.fill = pycopy.copy(cell.fill)
                    new_cell.border = pycopy.copy(cell.border)
                    new_cell.alignment = pycopy.copy(cell.alignment)
                    new_cell.number_format = cell.number_format

        for col_letter, dim in src_ws.column_dimensions.items():
            if dim.width:
                dst_ws.column_dimensions[col_letter].width = dim.width

        added_sheets.append(new_name)
        print(f"   ├─ Embedded color mapping sheet: [{new_name}] ({src_ws.max_row} rows)")

    return added_sheets


def _normalize_model_name(name):
    """Normalize a model name: uppercase, collapsed whitespace."""
    return re.sub(r"\s+", " ", str(name).strip().upper())


def load_model_cost_map(path, sheet_name=COST_SHEET_NAME):
    """
    Parse per-site B/C Model cost blocks from CoPQ workbook's "Current month" sheet.
    Headers: FTY | Model | Defect Qty | Ttl cost ($)
    Returns: dict of site -> list of (model_name, ttl_cost) tuples, sorted descending.
    """
    cost_map = {}

    if not os.path.exists(path):
        print(f"[!] Cost file not found: {path}")
        print("    Top 3 Models will fall back to Defect Qty ranking instead.")
        return cost_map

    try:
        wb_cost = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        print(f"[!] Could not read cost file ({path}): {e}")
        return cost_map

    if sheet_name not in wb_cost.sheetnames:
        print(f"[!] Sheet '{sheet_name}' not found in cost file. Available: {wb_cost.sheetnames}")
        return cost_map

    ws = wb_cost[sheet_name]
    max_r = ws.max_row or 1
    max_c = ws.max_column or 1

    for r in range(1, max_r + 1):
        for c in range(1, max_c - 2):
            v0 = ws.cell(row=r, column=c).value
            v1 = ws.cell(row=r, column=c + 1).value
            v2 = ws.cell(row=r, column=c + 2).value
            v3 = ws.cell(row=r, column=c + 3).value
            if not (
                v0 and str(v0).strip().lower() == "fty"
                and v1 and str(v1).strip().lower() == "model"
                and v2 and "defect qty" in str(v2).strip().lower()
                and v3 and str(v3).strip().lower().startswith("ttl cost")
            ):
                continue

            data_r = r + 1
            while data_r <= max_r:
                site_val = ws.cell(row=data_r, column=c).value
                model_val = ws.cell(row=data_r, column=c + 1).value
                cost_val = ws.cell(row=data_r, column=c + 3).value

                if not site_val and not model_val:
                    break

                if site_val and model_val and isinstance(cost_val, (int, float)):
                    site_key = str(site_val).strip().upper()
                    cost_map.setdefault(site_key, []).append(
                        (str(model_val).strip(), float(cost_val))
                    )

                data_r += 1

    for site_key in cost_map:
        cost_map[site_key].sort(key=lambda x: x[1], reverse=True)

    total_models = sum(len(v) for v in cost_map.values())
    print(f"  [+] Loaded B/C cost ranking for {len(cost_map)} site(s), {total_models} model rows (from: {path})")
    return cost_map


def find_matching_ftt_model(cost_model_name, available_models):
    """Match model name from cost file to FTT data model names."""
    norm_cost = _normalize_model_name(cost_model_name)

    # 1) Exact match
    for m in available_models:
        if _normalize_model_name(m) == norm_cost:
            return m

    # 2) Prefix match
    best_match = None
    best_overlap = 0
    for m in available_models:
        norm_m = _normalize_model_name(m)
        if norm_m.startswith(norm_cost) or norm_cost.startswith(norm_m):
            overlap = min(len(norm_m), len(norm_cost))
            if overlap > best_overlap:
                best_overlap = overlap
                best_match = m

    return best_match


def get_defect_color(defect_name, color_map, fallback_cache, unmatched_log=None):
    norm = _normalize_defect_name(defect_name)
    if norm in color_map:
        return color_map[norm]
    if norm.endswith("defect") and norm[:-len("defect")] in color_map:
        return color_map[norm[:-len("defect")]]
    if unmatched_log is not None:
        unmatched_log.add(str(defect_name))
    if norm in fallback_cache:
        return fallback_cache[norm]
    color = FALLBACK_PALETTE[len(fallback_cache) % len(FALLBACK_PALETTE)]
    fallback_cache[norm] = color
    return color


def get_label_font_color(hex_color):
    hex_color = hex_color.lstrip("#")
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    brightness = (r * 299 + g * 587 + b * 114) / 1000
    return "000000" if brightness >= 128 else "FFFFFF"


def _make_label_text_props(font_hex):
    char_props = CharacterProperties(solidFill=font_hex, sz=700, b=True)
    para_props = ParagraphProperties(defRPr=char_props)
    paragraph = Paragraph(pPr=para_props, endParaRPr=char_props)
    return RichText(bodyPr=RichTextProperties(), p=[paragraph])


def add_top3_chart(
    ws,
    header_row,
    first_model_row,
    last_model_row,
    defects_in_order,
    color_map,
    fallback_cache,
    anchor_cell,
    unmatched_log=None,
):
    """Add stacked-column 'Top 3 Defect' chart to summary worksheet."""
    if not defects_in_order or last_model_row < first_model_row:
        return

    num_defect_cols = len(defects_in_order)
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked"
    chart.overlap = 100
    chart.gapWidth = 150
    chart.title = "Top 3 Defect"
    chart.height = 7.5
    chart.width = 15.5
    chart.legend = None

    data = Reference(
        ws,
        min_col=2,
        max_col=1 + num_defect_cols,
        min_row=header_row,
        max_row=last_model_row,
    )
    cats = Reference(
        ws,
        min_col=1,
        min_row=first_model_row,
        max_row=last_model_row,
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    for series, defect_name in zip(chart.series, defects_in_order):
        fill_hex = get_defect_color(defect_name, color_map, fallback_cache, unmatched_log)
        series.graphicalProperties = GraphicalProperties(solidFill=fill_hex)
        font_hex = get_label_font_color(fill_hex)
        series.dLbls = DataLabelList(
            showVal=True,
            showLegendKey=False,
            showCatName=False,
            showSerName=False,
            showPercent=False,
            showBubbleSize=False,
            dLblPos="ctr",
            numFmt="0.00%",
            txPr=_make_label_text_props(font_hex),
        )

    chart.x_axis.delete = False
    chart.y_axis.delete = True
    chart.x_axis.majorGridlines = None
    chart.y_axis.majorGridlines = None

    ws.add_chart(chart, anchor_cell)


def write_custom_defect_legend(
    ws,
    site_name,
    header_row,
    defects_in_order,
    color_map,
    fallback_cache,
    key_col_idx,
    unmatched_log=None,
):
    """Write custom cell-based defect color key next to chart."""
    fill_site_header = PatternFill(start_color="1F4E78", fill_type="solid")
    font_site_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    key_border = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    font_key = Font(name="Calibri", size=9.5, color="000000")

    color_col_letter = get_column_letter(key_col_idx)
    text_col_letter = get_column_letter(key_col_idx + 1)

    ws.column_dimensions[color_col_letter].width = 5
    ws.column_dimensions[text_col_letter].width = 34

    ws.merge_cells(
        start_row=header_row,
        start_column=key_col_idx,
        end_row=header_row,
        end_column=key_col_idx + 1,
    )
    cell_hdr = ws.cell(row=header_row, column=key_col_idx, value=f"Site {site_name}")
    cell_hdr.fill = fill_site_header
    cell_hdr.font = font_site_header
    cell_hdr.alignment = Alignment(horizontal="center", vertical="center")
    cell_hdr.border = key_border

    cell_hdr_right = ws.cell(row=header_row, column=key_col_idx + 1)
    cell_hdr_right.border = key_border

    for i, defect_name in enumerate(defects_in_order, start=1):
        row_idx = header_row + i
        fill_hex = get_defect_color(defect_name, color_map, fallback_cache, unmatched_log)

        cell_color = ws.cell(row=row_idx, column=key_col_idx)
        cell_color.fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
        cell_color.border = key_border

        cell_text = ws.cell(row=row_idx, column=key_col_idx + 1, value=str(defect_name))
        cell_text.font = font_key
        cell_text.border = key_border
        cell_text.alignment = Alignment(horizontal="left", vertical="center")


# ----------------------------------------------------------------------
# 1. HELPER FUNCTIONS: SITE EXTRACTION & COLUMN MAPPING
# ----------------------------------------------------------------------
def auto_map_ftt_columns(df):
    df = df.loc[:, ~df.columns.duplicated()].copy()
    col_mapping = {}

    model2_col = None
    for col in df.columns:
        col_norm = re.sub(r"[\s_\-]+", "", str(col).strip().lower())
        if col_norm in ("model2", "model02", "modelno2", "modelno02"):
            model2_col = col
            break
    if model2_col is None:
        for col in df.columns:
            col_clean = str(col).strip().lower()
            if "model" in col_clean and "2" in col_clean:
                model2_col = col
                break

    if model2_col is not None:
        other_model_cols = [
            c for c in df.columns
            if c != model2_col and any(k in str(c).strip().lower() for k in ["model", "giày", "style", "art"])
        ]
        if other_model_cols:
            df = df.drop(columns=other_model_cols)
        col_mapping[model2_col] = "Model"

    for col in df.columns:
        if col == model2_col:
            continue

        col_clean = str(col).strip().lower()
        if any(k in col_clean for k in ["date", "ngày"]):
            col_mapping[col] = "Date"
        elif model2_col is None and any(k in col_clean for k in ["model", "giày", "style", "art"]):
            col_mapping[col] = "Model"
        elif any(k in col_clean for k in ["metric", "grade", "loại"]):
            col_mapping[col] = "Metric"
        elif any(
            k in col_clean
            for k in [
                "issues q'ty", "issues qty", "issue q'ty", "issue qty",
                "issues_qty", "ftt qty", "ftt_qty", "reject qty",
                "defect qty", "số lượng lỗi", "sl lỗi", "qty_defect",
            ]
        ):
            col_mapping[col] = "FTT_Qty"
        elif any(k in col_clean for k in ["defect", "issue", "mã lỗi", "tên lỗi", "lỗi"]) and "Defect_Issue" not in col_mapping.values():
            col_mapping[col] = "Defect_Issue"

    if "FTT_Qty" not in col_mapping.values():
        for col in df.columns:
            col_clean = str(col).strip().lower()
            if any(q in col_clean for q in ["q'ty", "qty", "số lượng"]):
                if not any(skip in col_clean for skip in ["total", "order", "inspected", "kiểm", "tổng", "plan", "check"]):
                    col_mapping[col] = "FTT_Qty"
                    break

    mapped_df = df.rename(columns=col_mapping)
    return mapped_df.loc[:, ~mapped_df.columns.duplicated()].copy()


def process_single_ftt_file(file_path):
    file_name = os.path.basename(file_path)
    site_name = extract_site_name(file_name)
    print(f"  [+] Reading file: {file_name} | Assigned Site: [{site_name}]")

    if file_path.endswith(".csv"):
        df_list = [pd.read_csv(file_path)]
    else:
        xls = pd.ExcelFile(file_path)
        df_list = [pd.read_excel(xls, sheet_name=sheet) for sheet in xls.sheet_names]

    full_df = pd.concat(df_list, ignore_index=True)
    df = auto_map_ftt_columns(full_df)

    for col in ["Model", "Date", "Metric", "Defect_Issue"]:
        if col not in df.columns:
            df[col] = np.nan
    if "FTT_Qty" not in df.columns:
        df["FTT_Qty"] = 0

    df["Site"] = site_name

    # Date normalization
    df["Date"] = pd.to_datetime(df["Date"], dayfirst=True, errors="coerce").dt.strftime("%Y-%m-%d")
    df["Date"] = df["Date"].fillna("N/A")

    # Filter Grade B/C defect criteria
    metric_str = df["Metric"].astype(str).str.strip()
    defect_str = df["Defect_Issue"].astype(str).str.strip()

    mask_metric = metric_str.str.contains(
        "B/Reject|C/Reject|B-Reject|C-Reject|^B$|^C$|Grade B|Grade C",
        case=False,
        na=False,
    )
    mask_defect = defect_str.str.contains("^B/|^C/|^B-|^C-", case=False, na=False)

    df = df[mask_metric | mask_defect].copy()
    df["FTT_Qty"] = pd.to_numeric(df["FTT_Qty"], errors="coerce").fillna(0).astype(int)

    print(f"     └─ Extracted {len(df)} B/C defect rows.")
    return df


# ----------------------------------------------------------------------
# 2. MAIN PROCESSING & DETAILED LOGGING
# ----------------------------------------------------------------------
# ----------------------------------------------------------------------
# 3. WORKBOOK CONSTRUCTION HELPERS (split out of generate_ftt_report)
# ----------------------------------------------------------------------
def _style_dataframe_sheet(ws, dataframe, output_cols):
    """Write a dataframe to a sheet with the standard FTT styling."""
    fill_site_header = PatternFill(start_color="1F4E78", fill_type="solid")
    font_regular = Font(name="Calibri", size=10, color="333333")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.append(output_cols)
    for c_idx in range(1, len(output_cols) + 1):
        cell = ws.cell(row=1, column=c_idx)
        cell.fill = fill_site_header
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, row in enumerate(dataframe.itertuples(index=False), start=2):
        ws.append(list(row))
        for c_idx in range(1, len(output_cols) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = font_regular
            cell.border = thin_border
            if c_idx in [1, 2, 4]:
                cell.alignment = Alignment(horizontal="center")
            elif c_idx == 6:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = "#,##0"

    # Fast vectorized column width sizing
    for col_idx, col_name in enumerate(output_cols, start=1):
        max_val_len = (dataframe[col_name].astype(str).str.len().max()
                       if not dataframe.empty else 0)
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(len(col_name), int(max_val_len or 0)) + 4

    ws.freeze_panes = "A2"


def build_workbook(combined_df, output_cols):
    """Create the report workbook: summary, Combined_All and Site_* sheets.
    Returns (wb, ws_pivot)."""
    wb = openpyxl.Workbook()
    ws_pivot = wb.active
    ws_pivot.title = "Top Defects Summary"
    ws_pivot.views.sheetView[0].showGridLines = True

    print("   ├─ Creating sheet: [Combined_All]")
    ws_combined = wb.create_sheet(title="Combined_All")
    ws_combined.views.sheetView[0].showGridLines = True
    _style_dataframe_sheet(ws_combined, combined_df[output_cols], output_cols)

    for site in ["VH", "VH2", "JV", "JV2"]:
        sheet_name = f"Site_{site}"
        ws_site = wb.create_sheet(title=sheet_name)
        ws_site.views.sheetView[0].showGridLines = True
        site_data = combined_df[combined_df["Site"] == site][output_cols].copy()
        _style_dataframe_sheet(ws_site, site_data, output_cols)
        print(f"   ├─ Created sheet: [{sheet_name}] ({len(site_data)} rows)")

    return wb, ws_pivot


def write_site_matrix(ws_pivot, combined_df, site, model_cost_map,
                      defect_color_map, fallback_color_cache,
                      unmatched_defects, current_row):
    """Write the 'Top 3 Models x Top 3 Defects' matrix block for one site.
    Returns the next free row, or None if the site had no data."""
    site_sheet = f"Site_{site}"
    site_df = combined_df[combined_df["Site"] == site]
    if site_df.empty:
        return None

    fill_site_header = PatternFill(start_color="1F4E78", fill_type="solid")
    fill_defect_header = PatternFill(start_color="D9E1F2", fill_type="solid")
    font_site_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_defect_header = Font(name="Calibri", size=10, bold=True, color="000000")
    font_model_name = Font(name="Calibri", size=10, bold=True, color="333333")
    font_regular = Font(name="Calibri", size=10, color="333333")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    available_models = site_df["Model"].dropna().unique().tolist()
    top_models = []
    for cost_model_name, ttl_cost in model_cost_map.get(site, []):
        matched = find_matching_ftt_model(cost_model_name, available_models)
        if matched is None:
            continue
        if matched not in top_models:
            top_models.append(matched)
        if len(top_models) == 3:
            break

    if not top_models:
        print(f"   [!] [{site}] No usable cost data matched - falling back to Defect Qty ranking.")
        top_models = (
            site_df.groupby("Model")["FTT_Qty"].sum().nlargest(3).index.tolist()
        )
    if not top_models:
        return None

    top_models_df = site_df[site_df["Model"].isin(top_models)]
    per_model_top_defects = {}
    selected_defects = set()
    for model_name in top_models:
        model_df = top_models_df[top_models_df["Model"] == model_name]
        model_top3 = (
            model_df.groupby("Defect_Issue")["FTT_Qty"].sum()
            .nlargest(3).index.tolist()
        )
        per_model_top_defects[model_name] = model_top3
        selected_defects.update(model_top3)

    defects_in_top_models = (
        top_models_df[top_models_df["Defect_Issue"].isin(selected_defects)]
        .groupby("Defect_Issue")["FTT_Qty"].sum()
        .sort_values(ascending=False).index.tolist()
    )

    header_row_idx = current_row

    # Site Header Row
    cell_site = ws_pivot.cell(row=current_row, column=1, value=site)
    cell_site.font = font_site_header
    cell_site.fill = fill_site_header
    cell_site.alignment = Alignment(horizontal="center", vertical="center")
    cell_site.border = thin_border

    # Defect Headers
    for d_idx, defect_name in enumerate(defects_in_top_models, start=2):
        cell_def = ws_pivot.cell(row=current_row, column=d_idx, value=defect_name)
        cell_def.font = font_defect_header
        cell_def.fill = fill_defect_header
        cell_def.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_def.border = thin_border

    ws_pivot.row_dimensions[current_row].height = 28
    current_row += 1
    first_model_row = current_row

    # Rows for Each Top Model
    for model_name in top_models:
        model_row = current_row
        cell_m = ws_pivot.cell(row=model_row, column=1, value=str(model_name))
        cell_m.font = font_model_name
        cell_m.border = thin_border
        cell_m.alignment = Alignment(vertical="center", wrap_text=True)
        ws_pivot.row_dimensions[model_row].height = 20

        model_top3_defects = per_model_top_defects[model_name]
        escaped_model_name = str(model_name).replace('"', '""')

        for d_idx, defect_name in enumerate(defects_in_top_models, start=2):
            cell_val = ws_pivot.cell(row=model_row, column=d_idx)
            cell_val.border = thin_border
            if defect_name not in model_top3_defects:
                continue

            defect_col_letter = get_column_letter(d_idx)
            formula = (
                f"=IFERROR("
                f"SUMIFS({site_sheet}!$F:$F, {site_sheet}!$C:$C, \"{escaped_model_name}\", {site_sheet}!$E:$E, {defect_col_letter}${header_row_idx})"
                f" / SUMIF({site_sheet}!$C:$C, \"{escaped_model_name}\", {site_sheet}!$F:$F), \"\")"
            )
            cell_val.value = formula
            cell_val.font = font_regular
            cell_val.alignment = Alignment(horizontal="right", vertical="center")
            cell_val.number_format = "0.00%"

        current_row += 1

    last_model_row = current_row - 1
    num_defect_cols = len(defects_in_top_models)

    # Stacked-column chart
    chart_col_idx = 2 + num_defect_cols + 1
    chart_anchor = f"{get_column_letter(chart_col_idx)}{header_row_idx}"
    add_top3_chart(
        ws_pivot, header_row_idx, first_model_row, last_model_row,
        defects_in_top_models, defect_color_map, fallback_color_cache,
        chart_anchor, unmatched_defects,
    )

    # Custom defect color key
    write_custom_defect_legend(
        ws_pivot, site, header_row_idx, defects_in_top_models,
        defect_color_map, fallback_color_cache, 21, unmatched_defects,
    )

    return current_row + 2


def save_report(wb, output_file):
    """Save the workbook, with a locked-file fallback. Returns True on success."""
    saved_path = save_with_fallback(wb, output_file)
    print("\n==================================================")
    print("[+] SUCCESS! FTT Report created with DR% matrix formulas & charts.")
    print(f"    Saved at: {saved_path}")
    print("==================================================\n")
    return True


def generate_ftt_report(
    input_folder=INPUT_FOLDER,
    output_file=OUTPUT_FILE,
    color_map_file=COLOR_MAP_FILE,
    cost_file=COST_FILE,
    cost_sheet_name=COST_SHEET_NAME,
    file_paths=None,
):
    print("\n==================================================")
    print("  STARTING AUTOMATED FTT DATA PROCESSING")
    print("==================================================")
    print(f"  Input Folder:   {input_folder}")
    print(f"  Output Report:  {output_file}")
    print(f"  Color Mapping:  {color_map_file}")
    print(f"  Cost Database:  {cost_file}")
    print("==================================================")

    if file_paths:
        # Explicit file selection (e.g. from the dashboard) overrides the folder scan.
        files = [f for f in file_paths
                 if os.path.splitext(f)[1].lower() in (".xlsx", ".xls", ".csv") and os.path.exists(f)]
        skipped = [f for f in file_paths if f not in files]
        for f in skipped:
            print(f"[-] Skipping missing/unsupported file: {f}")
        if not files:
            print("[-] Error: No valid data files among the selected FTT files!")
            return False
    else:
        if not os.path.exists(input_folder):
            print(f"[-] Error: Input directory not found: {input_folder}")
            return False
        output_dir = os.path.dirname(output_file)
        os.makedirs(output_dir, exist_ok=True)

        file_patterns = [
            os.path.join(input_folder, "*.xlsx"),
            os.path.join(input_folder, "*.xls"),
            os.path.join(input_folder, "*.csv"),
        ]

        files = []
        for pattern in file_patterns:
            files.extend(glob.glob(pattern))

        files = [f for f in files if not os.path.basename(f).startswith("~$")]

        if not files:
            print("[-] Error: No data files found in FTT input directory!")
            return False
    output_dir = os.path.dirname(output_file)
    os.makedirs(output_dir, exist_ok=True)

    print(f"[+] [Step 1/5] Found {len(files)} file(s) in input directory:")
    for f in files:
        print(f"   • {os.path.basename(f)}")

    print("\n[+] [Step 2/5] Parsing & Cleaning Data from files...")
    processed_dfs = []
    for f in files:
        try:
            df = process_single_ftt_file(f)
            processed_dfs.append(df)
        except Exception as e:
            print(f"[-] Error reading {os.path.basename(f)}: {e}")

    if not processed_dfs:
        print("[-] No valid data extracted!")
        return False

    combined_df = pd.concat(processed_dfs, ignore_index=True)
    print(f"\n[+] Total records combined: {len(combined_df)} rows across all sites.")

    # ------------------------------------------------------------------
    # BUILD EXCEL WORKBOOK
    # ------------------------------------------------------------------
    print("\n[+] [Step 3/5] Constructing Excel Workbook and Worksheets...")
    output_cols = ["Site", "Date", "Model", "Metric", "Defect_Issue", "FTT_Qty"]
    wb, ws_pivot = build_workbook(combined_df, output_cols)

    print("\n[+] [Step 4/5] Generating Matrix Style 'Top Defects Summary' Dashboard...")
    current_row = 2

    print("   ├─ Embedding Color_Defect.xlsx into report workbook...")
    embedded_color_sheets = embed_color_defect_workbook(wb, color_map_file)
    if embedded_color_sheets:
        defect_color_map = load_defect_color_map(wb, embedded_color_sheets)
    else:
        defect_color_map = load_defect_color_map(color_map_file)

    fallback_color_cache = {}
    unmatched_defects = set()

    print("   ├─ Loading model cost data (for Top 3 Model ranking)...")
    model_cost_map = load_model_cost_map(cost_file, cost_sheet_name)

    for site in ["VH", "VH2", "JV", "JV2"]:
        next_row = write_site_matrix(
            ws_pivot, combined_df, site, model_cost_map, defect_color_map,
            fallback_color_cache, unmatched_defects, current_row)
        if next_row is None:
            continue
        current_row = next_row

    # Column Dimensions
    ws_pivot.column_dimensions["A"].width = 35
    for col_idx in range(2, 25):
        col_let = get_column_letter(col_idx)
        ws_pivot.column_dimensions[col_let].width = 25

    # Save Output
    save_report(wb, output_file)
    return True


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Process FTT data and generate Top Defects matrix report.")
    parser.add_argument("--ftt-dir", default=INPUT_FOLDER, help="Path to FTT input directory")
    parser.add_argument("--output", default=OUTPUT_FILE, help="Path to output Excel report")
    parser.add_argument("--color-map", default=COLOR_MAP_FILE, help="Path to Color_Defect.xlsx")
    parser.add_argument("--cost-file", default=COST_FILE, help="Path to CoPQ_type_analysis.xlsx")
    parser.add_argument("--cost-sheet", default=COST_SHEET_NAME, help="Sheet name for cost data")
    args = parser.parse_args()

    generate_ftt_report(
        input_folder=args.ftt_dir,
        output_file=args.output,
        color_map_file=args.color_map,
        cost_file=args.cost_file,
        cost_sheet_name=args.cost_sheet,
    )